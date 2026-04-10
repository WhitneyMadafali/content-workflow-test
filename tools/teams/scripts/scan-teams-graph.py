#!/usr/bin/env python3
"""
Microsoft Teams scanner using Microsoft Graph API.

Features:
- Scan Team-backed groups and joined Teams
- Enumerate channels
- Export channel messages
- Export channel reply threads
- Optional filter for posts that include Microsoft Forms links (channel posts only; not the Forms API)
- Download channel files and include analysis metadata in JSON output
"""

import json
import re
import html
import sys
import argparse
import difflib
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional
from datetime import datetime

try:
    from msal import PublicClientApplication
    import requests
except ImportError:
    print("Required packages not installed.")
    print("Install with: pip install -r requirements.txt")
    sys.exit(1)

# Typical Microsoft Forms URLs embedded in Teams channel posts
_FORMS_HOST_RE = re.compile(r"forms\.(?:office|microsoft)\.com", re.IGNORECASE)

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from openpyxl import load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class TeamsScanner:
    AUTHORITY = "https://login.microsoftonline.com/common"
    GRAPH_ENDPOINT = "https://graph.microsoft.com/v1.0"

    # OVES custom Azure AD app registration
    CLIENT_ID = "1df26ef8-c7ce-4aee-aff2-bc36342362e0"

    SCOPES = [
        "User.Read",
        "Group.Read.All",
        "Team.ReadBasic.All",
        "Channel.ReadBasic.All",
        "ChannelMessage.Read.All",
        "Files.Read.All",
        "Sites.Read.All",
    ]

    def __init__(self, cache_file: str = ".teams_token_cache.json", workspace_root: Optional[Path] = None):
        script_dir = Path(__file__).parent.parent  # tools/teams
        self.cache_file = script_dir / cache_file
        self.token_cache = self._load_cache()
        self.app = PublicClientApplication(
            self.CLIENT_ID,
            authority=self.AUTHORITY,
            token_cache=self.token_cache
        )
        self.access_token: Optional[str] = None
        self.workspace_root = workspace_root or Path(__file__).parent.parent.parent.parent

    def _load_cache(self):
        from msal import SerializableTokenCache
        cache = SerializableTokenCache()
        if self.cache_file.exists():
            cache.deserialize(self.cache_file.read_text())
        return cache

    def _save_cache(self):
        if self.token_cache.has_state_changed:
            self.cache_file.write_text(self.token_cache.serialize())
            print(f"Token cache saved: {self.cache_file}")

    def authenticate(self) -> bool:
        print("\n" + "=" * 60)
        print("Microsoft Authentication Required")
        print("=" * 60)

        accounts = self.app.get_accounts()
        if accounts:
            result = self.app.acquire_token_silent(self.SCOPES, account=accounts[0])
            if result and "access_token" in result:
                self.access_token = result["access_token"]
                print(f"Using cached account: {accounts[0].get('username', 'unknown')}")
                return True

        print("Opening browser login...")
        result = self.app.acquire_token_interactive(scopes=self.SCOPES, prompt="select_account")
        if "access_token" in result:
            self.access_token = result["access_token"]
            self._save_cache()
            print("Authentication successful")
            return True

        print(f"Authentication failed: {result.get('error_description', result.get('error'))}")
        return False

    def _make_request(self, endpoint: str, method: str = "GET") -> Optional[Dict]:
        if not self.access_token:
            print("Not authenticated. Run authenticate() first.")
            return None

        url = endpoint if endpoint.startswith("http") else f"{self.GRAPH_ENDPOINT}{endpoint}"
        headers = {
            "Authorization": f"Bearer {self.access_token}",
            "Accept": "application/json",
        }

        try:
            response = requests.request(method, url, headers=headers, timeout=120)
            response.raise_for_status()
            return response.json()
        except requests.exceptions.HTTPError as e:
            print(f"HTTP error: {e}")
            print(f"Response: {response.text}")
            return None
        except Exception as e:
            print(f"Request failed: {e}")
            return None

    def _make_paginated_request(self, endpoint: str, max_items: Optional[int] = None) -> List[Dict]:
        items: List[Dict] = []
        next_endpoint: Optional[str] = endpoint

        while next_endpoint:
            payload = self._make_request(next_endpoint)
            if not payload:
                break

            page_items = payload.get("value", [])
            items.extend(page_items)
            if max_items is not None and len(items) >= max_items:
                return items[:max_items]

            next_link = payload.get("@odata.nextLink")
            if next_link and next_link.startswith(self.GRAPH_ENDPOINT):
                next_endpoint = next_link[len(self.GRAPH_ENDPOINT):]
            else:
                next_endpoint = next_link

        return items

    def _safe_path_part(self, value: str) -> str:
        cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", value or "")
        return cleaned.strip().strip(".") or "unknown"

    def _strip_html(self, value: str) -> str:
        if not value:
            return ""
        no_tags = re.sub(r"<[^>]+>", " ", value)
        return re.sub(r"\s+", " ", html.unescape(no_tags)).strip()

    def _sender_display_name(self, obj: Dict) -> Optional[str]:
        """Safely resolve sender display name from Graph message shape."""
        from_obj = obj.get("from") or {}
        user_obj = from_obj.get("user") or {}
        app_obj = from_obj.get("application") or {}
        return user_obj.get("displayName") or app_obj.get("displayName")

    def _normalize_for_match(self, value: Optional[str]) -> str:
        """Normalize text for tolerant matching across spacing/case/punctuation."""
        if not value:
            return ""
        out = []
        for char in value.lower():
            cat = unicodedata.category(char)
            if cat.startswith("L") or cat.startswith("N"):
                out.append(char)
            else:
                out.append(" ")
        return re.sub(r"\s+", " ", "".join(out)).strip()

    def _looks_like_match(self, value: Optional[str], needle: Optional[str], threshold: float = 0.82) -> bool:
        """
        Tolerant matcher:
        - case-insensitive
        - spacing-insensitive
        - punctuation-insensitive
        - approximate fuzzy fallback
        """
        if not needle:
            return True
        if not value:
            return False

        raw_h = value.lower()
        raw_n = needle.lower()
        if raw_n in raw_h:
            return True

        norm_h = self._normalize_for_match(value)
        norm_n = self._normalize_for_match(needle)
        if not norm_n:
            return True
        if norm_n in norm_h:
            return True

        compact_h = norm_h.replace(" ", "")
        compact_n = norm_n.replace(" ", "")
        if compact_n and compact_n in compact_h:
            return True

        # Fuzzy full-string check for short labels like team/channel names.
        if len(compact_h) <= 80:
            ratio = difflib.SequenceMatcher(None, compact_h, compact_n).ratio()
            if ratio >= threshold:
                return True

        # Fuzzy window check for longer text bodies.
        if len(compact_n) >= 4 and len(compact_h) > len(compact_n):
            win = len(compact_n)
            step = max(1, win // 3)
            start = 0
            while start < len(compact_h):
                segment = compact_h[start:start + win]
                if not segment:
                    break
                ratio = difflib.SequenceMatcher(None, segment, compact_n).ratio()
                if ratio >= threshold:
                    return True
                start += step

        return False

    def _text_has_forms_link(self, value: Optional[str]) -> bool:
        """True if text looks like it contains a Microsoft Forms URL."""
        if not value:
            return False
        return bool(_FORMS_HOST_RE.search(value))

    def _message_has_forms_link(self, message_obj: Dict) -> bool:
        """
        Detect Forms links in either visible message text or HTML hyperlink targets.
        Teams messages often store links in anchor href where visible text is generic
        (e.g. "Fill in form"), so checking stripped text alone is not enough.
        """
        body = (message_obj.get("body") or {})
        raw_content = body.get("content", "") or ""
        if self._text_has_forms_link(raw_content):
            return True
        visible_text = self._strip_html(raw_content)
        return self._text_has_forms_link(visible_text)

    def _contains_text(self, value: Optional[str], needle: Optional[str]) -> bool:
        return self._looks_like_match(value, needle)

    def list_teams_groups(self) -> List[Dict]:
        endpoint = "/groups?$select=id,displayName,description,mail,resourceProvisioningOptions&$top=999"
        groups = self._make_paginated_request(endpoint)
        result = []
        for group in groups:
            options = group.get("resourceProvisioningOptions", []) or []
            if "Team" in options:
                result.append({
                    "id": group.get("id"),
                    "displayName": group.get("displayName"),
                    "description": group.get("description"),
                    "mail": group.get("mail"),
                    "resourceProvisioningOptions": options,
                })
        return result

    def list_joined_teams(self) -> List[Dict]:
        endpoint = "/me/joinedTeams?$select=id,displayName,description"
        return self._make_paginated_request(endpoint)

    def list_team_channels(self, team_id: str) -> List[Dict]:
        endpoint = f"/teams/{team_id}/channels?$select=id,displayName,description,membershipType,webUrl"
        return self._make_paginated_request(endpoint)

    def list_channel_messages(self, team_id: str, channel_id: str, limit: int = 50) -> List[Dict]:
        safe_limit = max(1, min(limit, 200))
        endpoint = (
            f"/teams/{team_id}/channels/{channel_id}/messages"
            f"?$top={safe_limit}"
        )
        return self._make_paginated_request(endpoint, max_items=safe_limit)

    def list_channel_message_replies(self, team_id: str, channel_id: str, message_id: str, limit: int = 50) -> List[Dict]:
        safe_limit = max(1, min(limit, 200))
        endpoint = (
            f"/teams/{team_id}/channels/{channel_id}/messages/{message_id}/replies"
            f"?$top={safe_limit}"
        )
        return self._make_paginated_request(endpoint, max_items=safe_limit)

    def get_channel_files_folder(self, team_id: str, channel_id: str) -> Optional[Dict]:
        endpoint = f"/teams/{team_id}/channels/{channel_id}/filesFolder"
        return self._make_request(endpoint)

    def list_folder_contents(self, drive_id: str, item_id: str) -> List[Dict]:
        endpoint = f"/drives/{drive_id}/items/{item_id}/children"
        payload = self._make_request(endpoint)
        if payload and "value" in payload:
            return payload["value"]
        return []

    def download_file(self, drive_id: str, item_id: str, output_path: str) -> bool:
        endpoint = f"/drives/{drive_id}/items/{item_id}/content"
        url = f"{self.GRAPH_ENDPOINT}{endpoint}"
        headers = {"Authorization": f"Bearer {self.access_token}"}

        try:
            response = requests.get(url, headers=headers, stream=True, timeout=120)
            response.raise_for_status()
            with open(output_path, "wb") as out:
                for chunk in response.iter_content(chunk_size=8192):
                    out.write(chunk)
            return True
        except Exception as e:
            print(f"Download failed for {output_path}: {e}")
            return False

    def analyze_powerpoint(self, pptx_path: str) -> Dict:
        if not PPTX_AVAILABLE:
            return {"error": "python-pptx not installed", "install": "pip install python-pptx"}

        try:
            prs = Presentation(pptx_path)
            total_images = 0
            total_shapes = 0
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                row = {
                    "slide_number": i,
                    "layout": slide.slide_layout.name if hasattr(slide.slide_layout, "name") else "Unknown",
                    "shapes_count": len(slide.shapes),
                    "text_content": [],
                }
                for shape in slide.shapes:
                    total_shapes += 1
                    if hasattr(shape, "text") and shape.text.strip():
                        row["text_content"].append(shape.text.strip())
                    if hasattr(shape, "image"):
                        total_images += 1
                slides.append(row)

            return {
                "total_slides": len(prs.slides),
                "total_images": total_images,
                "total_shapes": total_shapes,
                "slides": slides,
            }
        except Exception as e:
            return {"error": f"PowerPoint analysis failed: {e}"}

    def analyze_excel(self, xlsx_path: str) -> Dict:
        if not OPENPYXL_AVAILABLE:
            return {"error": "openpyxl not installed", "install": "pip install openpyxl"}

        try:
            wb = load_workbook(xlsx_path, read_only=True, data_only=True)
            all_text = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                for row in sheet.iter_rows(values_only=True):
                    for cell in row:
                        if cell and isinstance(cell, str) and cell.strip():
                            all_text.append(cell.strip())
            terms = sorted({w for t in all_text for w in t.replace(",", " ").split() if len(w) > 1})
            return {
                "total_sheets": len(wb.sheetnames),
                "total_text_cells": len(all_text),
                "unique_terms_count": len(terms),
                "terms": terms[:100],
            }
        except Exception as e:
            return {"error": f"Excel analysis failed: {e}"}

    def analyze_downloaded_file(self, file_path: Path) -> Dict:
        ext = file_path.suffix.lower()
        if ext == ".pptx":
            return self.analyze_powerpoint(str(file_path))
        if ext == ".xlsx":
            return self.analyze_excel(str(file_path))
        if ext in [".txt", ".md", ".csv", ".json", ".xml", ".yml", ".yaml"]:
            try:
                preview = file_path.read_text(encoding="utf-8", errors="ignore")[:1500]
                return {"type": "text_preview", "preview": preview}
            except Exception as e:
                return {"type": "text_preview", "error": str(e)}
        return {"type": "metadata_only", "message": f"No parser configured for {ext}"}

    def scan_teams(self, output_file: Optional[str] = None, include_messages: bool = False, message_limit: int = 50,
                   include_replies: bool = True, reply_limit: int = 50, download_files: bool = False,
                   max_files_per_channel: int = 50, team_id: Optional[str] = None, team_name: Optional[str] = None,
                   channel_id: Optional[str] = None, channel_name: Optional[str] = None,
                   message_contains: Optional[str] = None, forms_links_only: bool = False) -> Dict:
        print("\n" + "=" * 60)
        print("Microsoft Teams Scanner")
        print("=" * 60)

        print("Fetching Team-backed groups...")
        teams_groups = self.list_teams_groups()
        print(f"Found {len(teams_groups)} Team-backed groups")

        print("Fetching joined Teams...")
        teams = self.list_joined_teams()
        if team_id:
            teams = [team for team in teams if (team.get("id") or "") == team_id]
        elif team_name:
            teams = [team for team in teams if self._contains_text(team.get("displayName"), team_name)]
        print(f"Found {len(teams)} joined Teams")

        group_by_id = {g.get("id"): g for g in teams_groups if g.get("id")}
        results = {
            "scan_date": datetime.now().isoformat(),
            "scope": {
                "include_messages": include_messages,
                "message_limit": message_limit,
                "include_replies": include_replies,
                "reply_limit": reply_limit,
                "download_files": download_files,
                "max_files_per_channel": max_files_per_channel,
                "team_id": team_id,
                "team_name": team_name,
                "channel_id": channel_id,
                "channel_name": channel_name,
                "message_contains": message_contains,
                "forms_links_only": forms_links_only,
            },
            "teams_group_count": len(teams_groups),
            "joined_team_count": len(teams),
            "teams_groups": teams_groups,
            "teams": [],
        }

        session_root = self.workspace_root / ".teams-session"
        if download_files:
            session_root.mkdir(parents=True, exist_ok=True)

        for i, team in enumerate(teams, 1):
            team_id = team.get("id")
            team_name = team.get("displayName", "Unknown Team")
            print(f"\n[{i}/{len(teams)}] Team: {team_name}")

            channels = self.list_team_channels(team_id) if team_id else []
            if channel_id:
                channels = [channel for channel in channels if (channel.get("id") or "") == channel_id]
            elif channel_name:
                channels = [channel for channel in channels if self._contains_text(channel.get("displayName"), channel_name)]
            print(f"  Channels: {len(channels)}")

            team_row = {
                "id": team_id,
                "displayName": team_name,
                "description": team.get("description"),
                "group": group_by_id.get(team_id),
                "channel_count": len(channels),
                "channels": [],
            }

            for j, channel in enumerate(channels, 1):
                channel_id = channel.get("id")
                channel_name = channel.get("displayName", "Unknown Channel")
                print(f"  [{j}/{len(channels)}] #{channel_name}")

                channel_row = {
                    "id": channel_id,
                    "displayName": channel_name,
                    "description": channel.get("description"),
                    "membershipType": channel.get("membershipType"),
                    "webUrl": channel.get("webUrl"),
                }

                if include_messages and team_id and channel_id:
                    messages = self.list_channel_messages(team_id, channel_id, message_limit)
                    total_replies = 0
                    msg_rows = []

                    for message in messages:
                        message_text = self._strip_html((message.get("body") or {}).get("content", ""))
                        if forms_links_only and not self._message_has_forms_link(message):
                            continue
                        if message_contains and not self._contains_text(message_text, message_contains):
                            # Keep replies out unless parent message matches requested text.
                            continue

                        msg_row = {
                            "id": message.get("id"),
                            "createdDateTime": message.get("createdDateTime"),
                            "lastModifiedDateTime": message.get("lastModifiedDateTime"),
                            "from": self._sender_display_name(message),
                            "subject": message.get("subject"),
                            "summary": message.get("summary"),
                            "body_preview": message_text[:500],
                            "webUrl": message.get("webUrl"),
                            "replyToId": message.get("replyToId"),
                        }

                        if include_replies and message.get("id"):
                            replies = self.list_channel_message_replies(team_id, channel_id, message["id"], reply_limit)
                            total_replies += len(replies)
                            msg_row["replies_count"] = len(replies)
                            msg_row["replies"] = [{
                                "id": reply.get("id"),
                                "createdDateTime": reply.get("createdDateTime"),
                                "lastModifiedDateTime": reply.get("lastModifiedDateTime"),
                                "from": self._sender_display_name(reply),
                                "subject": reply.get("subject"),
                                "summary": reply.get("summary"),
                                "body_preview": self._strip_html(reply.get("body", {}).get("content", ""))[:500],
                                "webUrl": reply.get("webUrl"),
                                "replyToId": reply.get("replyToId"),
                            } for reply in replies]

                        msg_rows.append(msg_row)

                    channel_row["messages_count"] = len(msg_rows)
                    channel_row["messages"] = msg_rows
                    if include_replies:
                        channel_row["replies_count"] = total_replies
                        print(f"    Messages: {len(msg_rows)}, Replies: {total_replies}")
                    else:
                        print(f"    Messages: {len(msg_rows)}")

                if download_files and team_id and channel_id:
                    folder_item = self.get_channel_files_folder(team_id, channel_id)
                    files_info: List[Dict] = []
                    if folder_item:
                        drive_id = folder_item.get("parentReference", {}).get("driveId")
                        item_id = folder_item.get("id")
                        children = self.list_folder_contents(drive_id, item_id) if drive_id and item_id else []
                        file_count = 0

                        for child in children:
                            if "file" not in child:
                                continue
                            if file_count >= max_files_per_channel:
                                break

                            file_name = child.get("name", "unknown")
                            local_dir = session_root / self._safe_path_part(team_name) / self._safe_path_part(channel_name)
                            local_dir.mkdir(parents=True, exist_ok=True)
                            local_path = local_dir / file_name

                            downloaded = local_path.exists() or self.download_file(drive_id, child.get("id"), str(local_path))
                            analysis = self.analyze_downloaded_file(local_path) if downloaded else {"error": "download_failed"}

                            files_info.append({
                                "name": file_name,
                                "size": child.get("size", 0),
                                "modified": child.get("lastModifiedDateTime"),
                                "created": child.get("createdDateTime"),
                                "mime_type": child.get("file", {}).get("mimeType"),
                                "web_url": child.get("webUrl"),
                                "session_file": str(local_path),
                                "downloaded": downloaded,
                                "analysis": analysis,
                            })
                            file_count += 1

                    channel_row["files_count"] = len(files_info)
                    channel_row["files"] = files_info
                    print(f"    Files captured: {len(files_info)}")

                team_row["channels"].append(channel_row)

            results["teams"].append(team_row)

        if output_file:
            output_path = Path(output_file)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(json.dumps(results, indent=2, ensure_ascii=False))
            print(f"\nSaved output: {output_path}")

        return results


def main():
    parser = argparse.ArgumentParser(description="Scan Teams groups/channels/messages/replies/files using Graph API")
    parser.add_argument("-o", "--output", help="Output JSON path", default=None)
    parser.add_argument("--list-teams", action="store_true", help="List joined Teams and exit")
    parser.add_argument("--include-channel-messages", action="store_true", help="Include channel messages")
    parser.add_argument("--message-limit", type=int, default=50, help="Max parent messages per channel (default: 50)")
    parser.add_argument("--skip-channel-replies", action="store_true", help="Skip channel reply threads")
    parser.add_argument("--reply-limit", type=int, default=50, help="Max replies per parent message (default: 50)")
    parser.add_argument("--download-channel-files", action="store_true", help="Download files from channel folders")
    parser.add_argument("--max-channel-files", type=int, default=50, help="Max files per channel (default: 50)")
    parser.add_argument("--team-id", help="Only scan this Team ID", default=None)
    parser.add_argument("--team-name", help="Only scan Teams matching this name text", default=None)
    parser.add_argument("--channel-id", help="Only scan this Channel ID", default=None)
    parser.add_argument("--channel-name", help="Only scan channels matching this name text", default=None)
    parser.add_argument("--message-contains", help="Only keep messages containing this text", default=None)
    parser.add_argument(
        "--forms-links-only",
        action="store_true",
        help="With --include-channel-messages: only keep posts whose body contains a Microsoft Forms URL "
        "(forms.office.com or forms.microsoft.com)",
    )

    args = parser.parse_args()

    scanner = TeamsScanner()
    if not scanner.authenticate():
        sys.exit(1)

    if args.list_teams:
        teams = scanner.list_joined_teams()
        if args.team_id:
            teams = [team for team in teams if (team.get("id") or "") == args.team_id]
        elif args.team_name:
            teams = [team for team in teams if scanner._contains_text(team.get("displayName"), args.team_name)]
        print(f"\nJoined teams: {len(teams)}")
        for team in teams:
            print(f"- {team.get('displayName')} ({team.get('id')})")
        return

    scanner.scan_teams(
        output_file=args.output,
        include_messages=args.include_channel_messages,
        message_limit=max(1, args.message_limit),
        include_replies=(args.include_channel_messages and not args.skip_channel_replies),
        reply_limit=max(1, args.reply_limit),
        download_files=args.download_channel_files,
        max_files_per_channel=max(1, args.max_channel_files),
        team_id=args.team_id,
        team_name=args.team_name,
        channel_id=args.channel_id,
        channel_name=args.channel_name,
        message_contains=args.message_contains,
        forms_links_only=args.forms_links_only,
    )


if __name__ == "__main__":
    main()
